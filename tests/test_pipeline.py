from __future__ import annotations

import io
import json
import zipfile
from types import SimpleNamespace

import cv2
import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import preprocess
from app.engines import TesseractEngine, claude_engine
from app.engines.claude_engine import PAGE_SCHEMA, ClaudeEngine
from app.exporters import export
from app.exporters.excel import as_number
from app.main import app
from app.schema import Block, Cell, ExtractedDocument, KeyValue, Page

from . import samples
from .benchmark import number_accuracy, output_lines, table_cell_accuracy


def sample_doc() -> ExtractedDocument:
    return ExtractedDocument(
        title="Order 7",
        pages=[Page(page_number=1, notes="corner torn", blocks=[
            Block(type="heading", text="Stationery order", level=1),
            Block(type="paragraph", text="Deliver by Friday.\nGate 2", uncertain=True),
            Block(type="list", items=["pens", "tape"], ordered=True),
            Block(type="key_value", pairs=[KeyValue(key="Phone", value="0301-555-0192"),
                                           KeyValue(key="Code", value="007", uncertain=True)]),
            Block(type="checkbox", text="Paid", checked=True),
            Block(type="table", header_rows=1, rows=[
                [Cell(text="Item"), Cell(text="Qty"), Cell(text="Price")],
                [Cell(text="Pencils"), Cell(text="12"), Cell(text="0.750")],
                [Cell(text="Ink"), Cell(text="1,200"), Cell(text="$3.10", uncertain=True)],
                [Cell(text="=SUM(A1)"), Cell(text="-4"), Cell(text="12.5%")],
                [Cell(text="short row")],
            ]),
            Block(type="signature", text="J. Carter"),
        ])],
    )


# ---------------- exporters ----------------

@pytest.mark.parametrize("text,expected", [
    ("12", (12, "0")),
    ("0.750", (0.75, "0.000")),  # trailing zero preserved through the number format
    ("1,200", (1200, "#,##0")),
    ("-4", (-4, "0")),
    ("$3.10", (3.1, '"$"0.00')),
    ("007", None),  # leading zeros are an ID, keep text
    ("4111111111111111", None),  # too long for Excel precision
    ("12 kg", None),
    ("3:30", None),
])
def test_number_detection(text, expected):
    got = as_number(text)
    if expected is None:
        assert got is None
    else:
        assert got[0] == pytest.approx(expected[0]) and got[1] == expected[1]


def test_xlsx_keeps_every_character():
    data, media, ext = export(sample_doc(), "xlsx")
    assert ext == "xlsx" and "spreadsheet" in media
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Table 1"]
    assert ws["A1"].value == "Item" and ws["A1"].font.bold
    assert ws["C2"].value == 0.75 and ws["C2"].number_format == "0.000"
    assert ws["B3"].value == 1200 and ws["B3"].number_format == "#,##0"
    assert ws["C3"].fill.fgColor.rgb.endswith("FFF2A8") and ws["C3"].comment is not None  # uncertain flagged
    assert ws["A4"].value == "=SUM(A1)" and ws["A4"].data_type == "s"  # never a live formula
    assert ws["C4"].value == pytest.approx(0.125) and ws["C4"].number_format.endswith("%")
    assert ws["C5"].value in (None, "")  # short row padded to rectangle
    content = wb["Content"]
    values = [c.value for row in content.iter_rows() for c in row]
    assert "0301-555-0192" in values and "007" in values and "Deliver by Friday." in values


def test_docx_contains_structure():
    data, _, _ = export(sample_doc(), "docx")
    d = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Stationery order" in text and "☑ Paid" in text and "Scanner note: corner torn" in text
    assert len(d.tables) == 2  # key/value + data table
    assert d.tables[1].cell(1, 2).text == "0.750"


def test_pdf_is_valid_and_survives_huge_cells():
    doc = sample_doc()
    doc.pages[0].blocks.append(Block(type="table", rows=[[Cell(text="word " * 3000), Cell(text="x")]]))
    data, media, _ = export(doc, "pdf")
    assert media == "application/pdf" and data.startswith(b"%PDF")


def test_csv_md_txt_json():
    doc = sample_doc()
    csv_bytes, media, ext = export(doc, "csv")
    assert ext == "csv" and "0.750" in csv_bytes.decode("utf-8-sig")
    doc.pages[0].blocks.append(Block(type="table", rows=[[Cell(text="a")]]))
    zip_bytes, media, ext = export(doc, "csv")
    assert ext == "zip" and len(zipfile.ZipFile(io.BytesIO(zip_bytes)).namelist()) == 2
    md = export(doc, "md")[0].decode()
    assert "| Item | Qty | Price |" in md and "$3.10 ⚠" in md
    assert "Gate 2" in export(doc, "txt")[0].decode()
    assert ExtractedDocument.model_validate_json(export(doc, "json")[0]).title == "Order 7"
    with pytest.raises(ValueError):
        export(doc, "exe")


# ---------------- preprocessing ----------------

def _text_page(angle: float = 0.0) -> np.ndarray:
    img = np.full((1400, 1100), 255, np.uint8)
    for i in range(14):
        cv2.putText(img, "The quick brown fox 12.50", (80, 120 + i * 85), cv2.FONT_HERSHEY_SIMPLEX, 1.4, 0, 3)
    return preprocess.rotate(img, angle) if angle else img


@pytest.mark.parametrize("angle", [-7.0, -2.0, 3.0, 9.0])
def test_deskew_recovers_angle(angle):
    est = preprocess.estimate_skew(_text_page(angle))
    assert est == pytest.approx(-angle, abs=0.8)


def test_flatten_finds_photographed_page():
    import random
    page = cv2.cvtColor(_text_page(), cv2.COLOR_GRAY2BGR)
    photo = samples.photo(page, random.Random(1), tilt=6, persp=0.05)
    flat, found = preprocess.flatten_page(photo)
    assert found
    h, w = flat.shape[:2]
    assert abs(h / w - 1400 / 1100) < 0.08  # aspect ratio of the real sheet restored
    assert flat.mean() > 200  # desk removed: mostly white paper


def test_prepare_upscales_and_keeps_small_dots():
    img = np.full((300, 400, 3), 255, np.uint8)
    cv2.circle(img, (200, 150), 1, (0, 0, 0), -1)  # a 3-pixel dot
    ok, buf = cv2.imencode(".png", img)
    p = preprocess.prepare(buf.tobytes())
    assert max(p.enhanced.shape) >= preprocess.MIN_LONG_SIDE
    assert (p.binary < 128).sum() > 20  # the dot survived enhancement + binarisation


def test_tiles_overlap_and_cover():
    img = np.zeros((900, 600), np.uint8)
    t = preprocess.tiles(img, 3)
    assert len(t) == 9 and t[0][0] == "top-left" and t[-1][0] == "bottom-right"
    assert all(tile.shape[0] > 300 and tile.shape[1] > 200 for _, tile in t)


def test_bad_image_rejected():
    with pytest.raises(ValueError):
        preprocess.prepare(b"not an image")


# ---------------- engines ----------------

@pytest.mark.skipif(not TesseractEngine().available(), reason="tesseract not installed")
def test_tesseract_reads_printed_invoice_exactly():
    s = samples.printed_invoice()
    page, _ = TesseractEngine().extract_page(preprocess.prepare(s.image), 1)
    doc = ExtractedDocument(pages=[page])
    assert table_cell_accuracy(s.table, doc) == 1.0
    assert number_accuracy(s, doc) == 1.0
    lines = output_lines(doc)
    assert all(l in lines for l in s.lines)


@pytest.mark.skipif(not TesseractEngine().available(), reason="tesseract not installed")
def test_tesseract_neat_handwriting():
    s = samples.neat_note()
    page, _ = TesseractEngine().extract_page(preprocess.prepare(s.image), 1)
    doc = ExtractedDocument(pages=[page])
    assert number_accuracy(s, doc) >= 0.8
    # "3:30 pm" must stay in running text, not become a bogus form field
    assert not any(b.type == "key_value" and any(kv.key.endswith("3") for kv in b.pairs) for b in page.blocks)


class FakeStream:
    def __init__(self, payload, calls, kwargs, stop_reason="end_turn"):
        self.payload, self.stop_reason = payload, stop_reason
        calls.append(kwargs)

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def get_final_message(self):
        return SimpleNamespace(stop_reason=self.stop_reason,
                               content=[SimpleNamespace(type="thinking"),
                                        SimpleNamespace(type="text", text=json.dumps(self.payload))])


CLAUDE_PAGE = {
    "title": "Stationery order", "language": "en", "handwritten": True, "notes": None,
    "blocks": [
        {"type": "heading", "text": "Stationery order", "level": 1, "items": None, "ordered": None, "rows": None,
         "header_rows": None, "pairs": None, "checked": None, "uncertain": False, "confidence": 0.98},
        {"type": "table", "text": None, "level": None, "items": None, "ordered": None,
         "rows": [[{"text": "Item", "uncertain": False}, {"text": "Price", "uncertain": False}],
                  [{"text": "Pencils", "uncertain": False}, {"text": "0.75", "uncertain": True}]],
         "header_rows": 1, "pairs": None, "checked": None, "uncertain": True, "confidence": 0.9},
    ],
}


@pytest.fixture
def fake_claude(monkeypatch):
    calls = []

    class FakeClient:
        def __init__(self, *a, **k):
            self.beta = SimpleNamespace(messages=SimpleNamespace(stream=lambda **kw: FakeStream(CLAUDE_PAGE, calls, kw)))

    monkeypatch.setattr(claude_engine.anthropic, "Anthropic", FakeClient)
    return calls


@pytest.mark.parametrize("precision,images,passes", [("fast", 1, 1), ("high", 6, 1), ("max", 11, 2)])
def test_claude_request_shape(fake_claude, precision, images, passes):
    s = samples.neat_note()
    page, meta = ClaudeEngine(api_key="test").extract_page(preprocess.prepare(s.image), 3, precision)
    assert len(fake_claude) == passes
    req = fake_claude[0]
    assert req["model"] == claude_engine.MODEL
    assert req["output_config"]["format"]["schema"] is PAGE_SCHEMA
    assert req["thinking"] == {"type": "adaptive"}
    assert req["fallbacks"] == "default"
    content = req["messages"][0]["content"]
    assert sum(1 for c in content if c["type"] == "image") == images
    if passes == 2:
        assert "DRAFT" in fake_claude[1]["messages"][0]["content"][-1]["text"]
    assert page.page_number == 3 and meta["handwritten"] is True
    assert page.blocks[1].rows[1][1].uncertain


def test_claude_refusal_is_reported(monkeypatch):
    calls = []

    class FakeClient:
        def __init__(self, *a, **k):
            self.beta = SimpleNamespace(messages=SimpleNamespace(
                stream=lambda **kw: FakeStream(CLAUDE_PAGE, calls, kw, stop_reason="refusal")))

    monkeypatch.setattr(claude_engine.anthropic, "Anthropic", FakeClient)
    from app.engines import EngineError
    with pytest.raises(EngineError, match="declined"):
        ClaudeEngine(api_key="k").extract_page(preprocess.prepare(samples.neat_note().image), 1, "fast")


def test_schema_is_strict():
    def walk(node):
        if isinstance(node, dict):
            if node.get("type") == "object":
                assert node["additionalProperties"] is False
                assert set(node["required"]) == set(node["properties"])
            for v in node.values():
                walk(v)
        elif isinstance(node, list):
            for v in node:
                walk(v)
    walk(PAGE_SCHEMA)


# ---------------- HTTP API ----------------

client = TestClient(app)


def test_index_and_engines():
    r = client.get("/")
    assert "OmniScan" in r.text
    assert r.headers["cross-origin-embedder-policy"] == "require-corp"
    e = client.get("/api/engines").json()
    assert {"claude", "tesseract"} <= set(e)


@pytest.mark.skipif(not TesseractEngine().available(), reason="tesseract not installed")
def test_extract_then_export_roundtrip():
    s = samples.printed_invoice()
    r = client.post("/api/extract", data={"engine": "tesseract", "precision": "fast"},
                    files=[("files", ("a.png", s.image, "image/png")), ("files", ("b.png", s.image, "image/png"))])
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["engine"] == "tesseract" and len(body["document"]["pages"]) == 2
    assert body["pages"][0]["enhanced_preview"].startswith("data:image/jpeg;base64,")
    for fmt in ("xlsx", "docx", "pdf", "csv", "json", "txt", "md"):
        x = client.post("/api/export", json={"document": body["document"], "format": fmt, "filename": "inv/../x"})
        assert x.status_code == 200, (fmt, x.text)
        assert 'filename="inv x.' in x.headers["content-disposition"] or 'filename="invx.' in x.headers["content-disposition"]


def test_extract_via_claude_with_browser_key(fake_claude):
    s = samples.neat_note()
    r = client.post("/api/extract", data={"engine": "auto", "precision": "high"},
                    headers={"X-Anthropic-Key": "sk-test"}, files=[("files", ("n.png", s.image, "image/png"))])
    assert r.status_code == 200, r.text
    assert r.json()["engine"] == "claude" and r.json()["document"]["title"] == "Stationery order"


def test_extract_errors(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_AUTH_TOKEN", raising=False)
    img = samples.neat_note().image
    r = client.post("/api/extract", data={"engine": "claude"}, files=[("files", ("n.png", img, "image/png"))])
    assert r.status_code == 400 and "API key" in r.json()["detail"]
    r = client.post("/api/extract", data={"engine": "tesseract"}, files=[("files", ("n.txt", b"hello", "text/plain"))])
    assert r.status_code == 400
    r = client.post("/api/extract", data={"precision": "ultra"}, files=[("files", ("n.png", img, "image/png"))])
    assert r.status_code == 400
