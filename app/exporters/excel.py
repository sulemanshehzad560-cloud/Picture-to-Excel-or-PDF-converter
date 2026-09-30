from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..schema import ExtractedDocument, block_to_text

HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(bold=True, color="FFFFFF")
UNCERTAIN_FILL = PatternFill("solid", fgColor="FFF2A8")
THIN = Side(style="thin", color="B7BDC9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

# Only unambiguous numbers become numeric cells; anything else stays text exactly as read.
NUM_RE = re.compile(r"^(?P<cur>[$€£¥₹])?\s?(?P<neg>-)?(?P<int>\d{1,3}(?:,\d{3})+|\d+)(?:\.(?P<dec>\d+))?(?P<pct>%)?$")


def as_number(text: str):
    """Return (value, number_format) or None. Keeps leading zeros ('007') and IDs as text."""
    s = text.strip()
    m = NUM_RE.match(s)
    if not m:
        return None
    int_part = m.group("int")
    if len(int_part) > 1 and int_part.startswith("0"):
        return None
    if len(int_part.replace(",", "")) > 15:  # beyond Excel precision (card/account numbers)
        return None
    dec = m.group("dec") or ""
    value = float(int_part.replace(",", "") + ("." + dec if dec else ""))
    if m.group("neg"):
        value = -value
    fmt = ("#,##0" if "," in int_part else "0") + ("." + "0" * len(dec) if dec else "")
    if m.group("pct"):
        return value / 100, fmt + "%"
    if m.group("cur"):
        fmt = f'"{m.group("cur")}"{fmt}'
    return (int(value) if not dec and not m.group("pct") else value), fmt


def _safe_title(name: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "-", name)[:31] or "Sheet"
    base, n = name, 2
    while name in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def _write_cell(ws, row: int, col: int, text: str, uncertain: bool, header: bool = False, struck: bool = False):
    cell = ws.cell(row=row, column=col)
    parsed = None if header else as_number(text)
    if parsed:
        cell.value, cell.number_format = parsed
    else:
        cell.value = text
        if text.startswith("="):
            cell.data_type = "s"  # handwritten "=5+3" is text, never a live formula
    cell.border = BORDER
    cell.alignment = WRAP
    if header:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    if uncertain:
        cell.fill = UNCERTAIN_FILL
        cell.comment = Comment("Low-confidence reading - please verify against the original.", "Scanner")
    if struck:
        cell.font = Font(strike=True, color="7A7F8C", bold=header)
    return cell


def _autofit(ws):
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            longest = max((len(part) for part in str(c.value).split("\n")), default=0)
            widths[c.column] = max(widths.get(c.column, 0), longest)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(60, max(8, w + 2))


def to_xlsx(doc: ExtractedDocument) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    tables = doc.tables()

    for page, n, block in tables:
        title = f"Table {n}" if len(doc.pages) == 1 else f"P{page} Table {n}"
        ws = wb.create_sheet(_safe_title(title, used))
        hdr = block.header_rows or 0
        for r, row in enumerate(block.rows, 1):
            for c, cell in enumerate(row, 1):
                _write_cell(ws, r, c, cell.text, cell.uncertain, header=r <= hdr, struck=cell.struck)
        if hdr:
            ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
        _autofit(ws)

    # Everything that isn't a table: forms get Field/Value columns, the rest is one row per line.
    ws = wb.create_sheet(_safe_title("Content", used), 0 if not tables else None)
    r = 1
    for c, h in enumerate(["Page", "Type", "Field / Text", "Value"], 1):
        _write_cell(ws, r, c, h, False, header=True)
    for page in doc.pages:
        for b in page.blocks:
            if b.type == "table":
                continue
            if b.type == "key_value":
                for kv in b.pairs or []:
                    r += 1
                    _write_cell(ws, r, 1, str(page.page_number), False)
                    _write_cell(ws, r, 2, "field", False)
                    _write_cell(ws, r, 3, kv.key, kv.uncertain)
                    _write_cell(ws, r, 4, kv.value, kv.uncertain)
                continue
            if b.type == "checkbox":
                r += 1
                _write_cell(ws, r, 1, str(page.page_number), False)
                _write_cell(ws, r, 2, "checkbox", False)
                _write_cell(ws, r, 3, b.text or "", b.uncertain)
                _write_cell(ws, r, 4, "☑ checked" if b.checked else "☐ unchecked", b.uncertain)
                continue
            lines = (b.items or []) if b.type == "list" else block_to_text(b).split("\n")
            for line in lines:
                r += 1
                _write_cell(ws, r, 1, str(page.page_number), False)
                _write_cell(ws, r, 2, b.type, False)
                _write_cell(ws, r, 3, line, b.uncertain)
    ws.freeze_panes = "A2"
    _autofit(ws)
    if r == 1 and tables:
        wb.remove(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
