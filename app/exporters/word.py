from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..schema import Block, ExtractedDocument


def _run(par, text: str, uncertain: bool = False, bold: bool = False):
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = par.add_run(line)
        run.bold = bold
        if uncertain:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if i < len(lines) - 1:
            run.add_break()


def _shade(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _table(document, block: Block):
    rows = block.rows or []
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    hdr = block.header_rows or 0
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            target = table.cell(r, c)
            target.paragraphs[0].text = ""
            _run(target.paragraphs[0], cell.text, cell.uncertain, bold=r < hdr)
            if r < hdr:
                _shade(target, "DCE3F0")
    document.add_paragraph()


def to_docx(doc: ExtractedDocument) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    if doc.title:
        document.add_heading(doc.title, level=0)

    for p_i, page in enumerate(doc.pages):
        if p_i:
            document.add_page_break()
        for b in page.blocks:
            if b.type == "heading":
                h = document.add_heading(level=min(3, max(1, b.level or 1)))
                _run(h, b.text or "", b.uncertain)
            elif b.type == "paragraph":
                _run(document.add_paragraph(), b.text or "", b.uncertain)
            elif b.type == "list":
                style = "List Number" if b.ordered else "List Bullet"
                for item in b.items or []:
                    _run(document.add_paragraph(style=style), item, b.uncertain)
            elif b.type == "table":
                _table(document, b)
            elif b.type == "key_value":
                pairs = b.pairs or []
                if pairs:
                    t = document.add_table(rows=len(pairs), cols=2)
                    t.style = "Table Grid"
                    for i, kv in enumerate(pairs):
                        _run(t.cell(i, 0).paragraphs[0], kv.key, kv.uncertain, bold=True)
                        _run(t.cell(i, 1).paragraphs[0], kv.value, kv.uncertain)
                        _shade(t.cell(i, 0), "EEF1F7")
                    document.add_paragraph()
            elif b.type == "checkbox":
                _run(document.add_paragraph(), ("☑ " if b.checked else "☐ ") + (b.text or ""), b.uncertain)
            elif b.type in ("signature", "figure"):
                par = document.add_paragraph()
                run = par.add_run(("Signature: " if b.type == "signature" else "") + (b.text or ""))
                run.italic = True
                run.font.color.rgb = RGBColor(0x66, 0x6E, 0x80)
        if page.notes:
            par = document.add_paragraph()
            run = par.add_run(f"Scanner note: {page.notes}")
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
